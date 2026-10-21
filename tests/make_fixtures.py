"""Generate test files for MCH PDF Organizer.  Usage: python tests/make_fixtures.py <outdir>

Every page of every text document carries a unique marker (e.g. "MARK-DOCXA-P2") so
tests can read the exported PDF and verify the exact page order.
"""
import io
import os
import shutil
import subprocess
import sys

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Border, Font, PatternFill, Side
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader, PdfWriter
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas

OUT = sys.argv[1] if len(sys.argv) > 1 else "fixtures"
os.makedirs(OUT, exist_ok=True)
P = lambda n: os.path.join(OUT, n)


def font(size):
    for f in ("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",):
        if os.path.exists(f):
            return ImageFont.truetype(f, size)
    return ImageFont.load_default()


def picture(w, h, label, color, alpha=False):
    im = Image.new("RGBA" if alpha else "RGB", (w, h), color + ((160,) if alpha else ()))
    d = ImageDraw.Draw(im)
    d.rectangle([4, 4, w - 5, h - 5], outline=(0, 0, 0), width=max(4, w // 200))
    d.line([0, 0, w, h], fill=(255, 255, 255), width=max(3, w // 300))
    d.ellipse([w * .35, h * .35, w * .65, h * .65], outline=(255, 255, 255), width=max(3, w // 200))
    d.text((w * .05, h * .05), label, fill=(255, 255, 255), font=font(max(20, w // 18)))
    return im


thin = Side(style="thin", color="000000")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = PatternFill("solid", fgColor="0F5257")


def excel_wide():
    wb = Workbook()
    ws = wb.active
    ws.title = "Audit"
    ws["A1"] = "MARK-XLSWIDE-S1"
    ws["A1"].font = Font(bold=True, size=14)
    cols = 28
    for c in range(1, cols + 1):
        cell = ws.cell(row=3, column=c, value=f"Criterion {c}")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.border = border
        ws.column_dimensions[cell.column_letter].width = 14
    for r in range(4, 64):
        for c in range(1, cols + 1):
            cell = ws.cell(row=r, column=c, value=(r * c) % 3 if c > 1 else f"Unit {r-3}")
            cell.border = border
    ws.cell(row=3, column=cols, value="LASTCOL-MARK")
    s2 = wb.create_sheet("Notes")
    s2["A1"] = "MARK-XLSWIDE-S2"
    s2["A2"] = "Second worksheet, narrow content."
    s3 = wb.create_sheet("Lists")
    s3["A1"] = "MARK-XLSWIDE-HIDDEN"
    s3.sheet_state = "hidden"
    wb.save(P("audit_wide.xlsx"))


def excel_narrow():
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "MARK-XLSNARROW-S1"
    for i, h in enumerate(["Unit", "Score", "Findings", "Status"], 1):
        c = ws.cell(row=2, column=i, value=h)
        c.font = Font(bold=True)
        c.border = border
    for r in range(3, 13):
        for i in range(1, 5):
            ws.cell(row=r, column=i, value=f"R{r}C{i}").border = border
    buf = io.BytesIO()
    picture(400, 200, "XL IMG", (184, 150, 63)).save(buf, "PNG")
    buf.seek(0)
    img = XLImage(buf)
    ws.add_image(img, "F2")
    wb.save(P("summary_narrow.xlsx"))


def word_docx():
    doc = Document()
    doc.add_heading("MARK-DOCXA-P1", level=1)
    doc.add_paragraph("Nursing quality report. ")
    p = doc.add_paragraph()
    p.add_run("تقرير الجودة التمريضية - تجمع المدينة المنورة الصحي").font.size = Pt(14)
    t = doc.add_table(rows=4, cols=3)
    t.style = "Table Grid"
    for r in range(4):
        for c in range(3):
            t.cell(r, c).text = f"Cell {r},{c}"
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("MARK-DOCXA-P2", level=1)
    buf = io.BytesIO()
    picture(1200, 600, "WIDE 2:1", (20, 107, 113)).save(buf, "PNG")
    buf.seek(0)
    doc.add_picture(buf, width=Inches(6))
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("MARK-DOCXA-P3", level=1)
    doc.add_paragraph("Final page.")
    doc.save(P("report_a.docx"))

    doc = Document()
    doc.add_heading("MARK-DOCXB-P1", level=1)
    doc.add_paragraph("Policy document, page one.")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("MARK-DOCXB-P2", level=1)
    doc.save(P("policy_b.docx"))


def legacy(src, ext, final):
    tmp = P("_legacy")
    os.makedirs(tmp, exist_ok=True)
    subprocess.run(["soffice", "--headless", "--convert-to", ext, "--outdir", tmp, src],
                   check=True, capture_output=True, timeout=180)
    produced = os.path.join(tmp, os.path.splitext(os.path.basename(src))[0] + "." + ext)
    shutil.move(produced, P(final))
    shutil.rmtree(tmp)


def pdfs():
    c = canvas.Canvas(P("scan_c.pdf"), pagesize=A4)
    for i in range(1, 5):
        if i == 3:
            c.setPageSize(landscape(A4))
        else:
            c.setPageSize(A4)
        c.setFont("Helvetica-Bold", 28)
        c.drawString(72, 500, f"MARK-PDFC-P{i}")
        c.rect(40, 40, 200, 100)
        c.showPage()
    c.save()

    c = canvas.Canvas(P("big_150.pdf"), pagesize=A4)
    for i in range(1, 151):
        c.setFont("Helvetica", 20)
        c.drawString(72, 700, f"MARK-BIG-P{i}")
        c.showPage()
    c.save()

    w = PdfWriter()
    w.append(P("scan_c.pdf"))
    w.encrypt(user_password="secret", owner_password="owner")
    with open(P("locked.pdf"), "wb") as f:
        w.write(f)

    data = open(P("scan_c.pdf"), "rb").read()
    open(P("broken.pdf"), "wb").write(data[: len(data) // 3])


def images():
    picture(1500, 2250, "PORTRAIT", (90, 60, 120)).save(P("photo_portrait.jpg"), quality=92)
    picture(2400, 1200, "LANDSCAPE 2:1", (20, 107, 113), alpha=True).save(P("chart_landscape.png"))
    # EXIF orientation 6 = rotate 90 CW when shown. Stored 1600x1000 -> displays 1000x1600 (portrait)
    im = picture(1600, 1000, "EXIF6", (160, 80, 40))
    exif = Image.Exif()
    exif[0x0112] = 6
    im.save(P("phone_exif.jpeg"), quality=90, exif=exif)
    picture(160, 90, "S", (40, 40, 40)).save(P("tiny.png"))
    # large photo
    import random
    big = picture(6000, 4000, "LARGE", (30, 80, 60))
    d = ImageDraw.Draw(big)
    rnd = random.Random(1)
    for _ in range(6000):
        x, y = rnd.randrange(6000), rnd.randrange(4000)
        d.rectangle([x, y, x + 20, y + 20], fill=(rnd.randrange(256), rnd.randrange(256), rnd.randrange(256)))
    big.save(P("large_photo.jpg"), quality=95)


def invalid():
    open(P("not_really.docx"), "w").write("plain text pretending to be a Word file")
    open(P("notes.txt"), "w").write("unsupported")
    open(P("empty.pdf"), "wb").close()
    open(P("fake_image.png"), "wb").write(b"\x89PNG\r\n\x1a\n" + b"garbage" * 50)
    # password-protected Word file
    try:
        import msoffcrypto
        from msoffcrypto.format.ooxml import OOXMLFile
        with open(P("policy_b.docx"), "rb") as fin, open(P("locked.docx"), "wb") as fout:
            OOXMLFile(fin).encrypt("secret", fout)
    except Exception as e:
        print("could not create locked.docx:", e)


excel_wide()
excel_narrow()
word_docx()
pdfs()
images()
invalid()
legacy(P("policy_b.docx"), "doc", "policy_legacy.doc")
legacy(P("summary_narrow.xlsx"), "xls", "summary_legacy.xls")
print("fixtures in", OUT)
for n in sorted(os.listdir(OUT)):
    print(f"  {n:24s} {os.path.getsize(P(n)):>10,d}")
